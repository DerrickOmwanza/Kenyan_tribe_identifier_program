"""Ingest a labelled name dataset (.xlsx) into the corpus store."""

from __future__ import annotations

import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from .. import __version__
from ..hashing import sha256_file
from .build import build
from .store import connect

POSITIONS = ("fname", "mname", "sname")


class IngestError(RuntimeError):
    pass


def _resolve_columns(header: list, wanted: dict[str, str]) -> dict[str, int]:
    """Map logical field -> column index, matching header names case-insensitively."""
    lookup = {str(name).strip().lower(): idx for idx, name in enumerate(header) if name is not None}
    resolved: dict[str, int] = {}
    missing: list[str] = []
    for field, column_name in wanted.items():
        idx = lookup.get(column_name.strip().lower())
        if idx is None:
            missing.append(column_name)
        else:
            resolved[field] = idx
    if missing:
        raise IngestError(
            f"input is missing required column(s): {', '.join(missing)}; "
            f"header was: {', '.join(str(h) for h in header)}"
        )
    return resolved


def _read_xlsx_rows(path: Path, sheet: str | None = None):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is None:
            ws = wb[wb.sheetnames[0]]
        elif sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            raise IngestError(
                f"sheet {sheet!r} not found; available: {', '.join(wb.sheetnames)}"
            )
        for row in ws.iter_rows(values_only=True):
            yield row
    finally:
        wb.close()


def _verify_hash(path: Path, manifest_path: Path | None, force: bool) -> tuple[str, bool]:
    actual = sha256_file(path).upper()
    if manifest_path and manifest_path.is_file():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        expected = str(manifest.get("sha256", "")).upper()
        if expected and expected != actual:
            message = (
                f"corpus file hash mismatch\n  file:     {path}\n"
                f"  expected: {expected}\n  actual:   {actual}\n"
                f"  manifest: {manifest_path}"
            )
            if not force:
                raise IngestError(message + "\nRe-run with --force to ingest anyway.")
        return actual, bool(expected) and expected == actual
    return actual, False


def ingest(
    input_path: str | Path,
    db_path: str | Path,
    *,
    source_id: str = "ref_v1",
    manifest_path: str | Path | None = None,
    fname_col: str = "fname",
    mname_col: str = "mname",
    sname_col: str = "sname",
    tribe_col: str = "Tribe",
    force: bool = False,
) -> dict[str, object]:
    input_path = Path(input_path)
    if not input_path.is_file():
        raise IngestError(f"corpus file not found: {input_path}")
    if input_path.suffix.lower() != ".xlsx":
        raise IngestError(f"expected an .xlsx file, got: {input_path.name}")

    sha_actual, verified = _verify_hash(
        input_path, Path(manifest_path) if manifest_path else None, force
    )

    rows = _read_xlsx_rows(input_path)
    header = list(next(rows))
    cols = _resolve_columns(
        header,
        {"fname": fname_col, "mname": mname_col, "sname": sname_col, "tribe": tribe_col},
    )

    def cell(row, idx: int) -> str:
        value = row[idx] if idx < len(row) else None
        return "" if value is None else str(value).strip()

    records: list[tuple[int, str | None, str | None, str | None, str]] = []
    skipped_blank_tribe = 0
    for row_index, row in enumerate(rows, start=1):
        tribe = cell(row, cols["tribe"])
        if not tribe:
            skipped_blank_tribe += 1
            continue
        raw = {p: cell(row, cols[p]) for p in POSITIONS}
        records.append(
            (row_index, raw["fname"] or None, raw["mname"] or None, raw["sname"] or None, tribe)
        )

    if not records:
        raise IngestError("no rows with a non-empty tribe label were found")

    conn = connect(db_path)
    try:
        _replace_source(conn, source_id, input_path.name, sha_actual, len(records), verified)
        conn.executemany(
            "INSERT INTO records(source_id, row_index, fname_raw, mname_raw, sname_raw, tribe_raw)"
            " VALUES (?, ?, ?, ?, ?, ?)",
            [(source_id, *rec) for rec in records],
        )
        summary = build(conn)
        conn.commit()
    finally:
        conn.close()

    return {
        "source_id": source_id,
        "file": input_path.name,
        "sha256": sha_actual,
        "hash_verified": verified,
        "records_ingested": len(records),
        "skipped_blank_tribe": skipped_blank_tribe,
        **summary,
    }


def _replace_source(
    conn: sqlite3.Connection, source_id: str, filename: str, sha256: str, row_count: int, verified: bool
) -> None:
    conn.execute("DELETE FROM sources WHERE source_id = ?", (source_id,))  # cascades to records
    conn.execute(
        "INSERT INTO sources"
        "(source_id, filename, sha256, row_count, ingested_at, tool_version, hash_verified)"
        " VALUES (?, ?, ?, ?, ?, ?, ?)",
        (source_id, filename, sha256, row_count,
         datetime.now(timezone.utc).isoformat(), __version__, int(verified)),
    )
