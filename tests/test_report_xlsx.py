"""Tests for the formatted .xlsx writer."""

from __future__ import annotations

import openpyxl
import pytest

from kne.pipeline import ClassifyResult
from kne.report_xlsx import write_workbook

FIELDNAMES = ["fname", "sname", "tribe_predicted", "tribe_confidence", "tribe_status",
              "tribe_basis", "tribe_alt_candidates"]


def _row(fname, sname, tribe, conf, status):
    return {"fname": fname, "sname": sname, "tribe_predicted": tribe,
            "tribe_confidence": f"{conf:.4f}", "tribe_status": status,
            "tribe_basis": "sname:X", "tribe_alt_candidates": ""}


def _result(rows):
    review = [r for r in rows if r["tribe_status"] in {"review", "conflict", "generic", "unverified"}]
    counts: dict[str, int] = {}
    for r in rows:
        counts[r["tribe_status"]] = counts.get(r["tribe_status"], 0) + 1
    return ClassifyResult(
        fieldnames=FIELDNAMES, rows=rows, review_rows=review, status_counts=counts,
        mode="fresh",
        manifest={"input_file": "in.csv", "corpus_version": "abc123", "mode": "fresh",
                  "status_counts": counts, "gate": {"tau": 0.75}},
    )


@pytest.fixture()
def workbook(tmp_path):
    rows = [
        _row("JOHN", "MWANGI", "Agikuyu (Kikuyu)", 0.98, "assigned"),
        _row("MARY", "OTIENO", "Luo", 0.95, "assigned"),
        _row("PETER", "JUMA", "", 0.44, "review"),
        _row("JANE", "MUTUA", "Akamba (Kamba)", 0.91, "conflict"),
    ]
    path = tmp_path / "wb.xlsx"
    write_workbook(path, _result(rows))
    return openpyxl.load_workbook(path)


def test_sheet_layout(workbook):
    assert workbook.sheetnames == ["Classified Data", "Summary", "Review Queue", "Run Manifest"]


def test_classified_data_sheet(workbook):
    ws = workbook["Classified Data"]
    assert ws.max_row == 5  # header + 4
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    header = [c.value for c in ws[1]]
    conf_col = header.index("tribe_confidence") + 1
    assert isinstance(ws.cell(row=2, column=conf_col).value, float)


def test_review_queue_only_reviewish(workbook):
    ws = workbook["Review Queue"]
    header = [c.value for c in ws[1]]
    status_col = header.index("tribe_status") + 1
    statuses = {ws.cell(row=r, column=status_col).value for r in range(2, ws.max_row + 1)}
    assert statuses <= {"review", "conflict", "generic", "unverified"}
    assert statuses == {"review", "conflict"}


def test_review_queue_sorted_by_confidence(workbook):
    ws = workbook["Review Queue"]
    header = [c.value for c in ws[1]]
    conf_col = header.index("tribe_confidence") + 1
    values = [ws.cell(row=r, column=conf_col).value for r in range(2, ws.max_row + 1)]
    assert values == sorted(values)


def test_summary_and_manifest_sheets(workbook):
    summary_text = "\n".join(
        str(c.value) for row in workbook["Summary"].iter_rows() for c in row if c.value
    )
    assert "Total rows" in summary_text and "assigned" in summary_text

    manifest_rows = {
        row[0].value: row[1].value for row in workbook["Run Manifest"].iter_rows() if row[0].value
    }
    assert manifest_rows["corpus_version"] == "abc123"
    assert manifest_rows["gate.tau"] == "0.75"


def test_no_review_sheet_when_empty(tmp_path):
    rows = [_row("JOHN", "MWANGI", "Agikuyu (Kikuyu)", 0.98, "assigned")]
    path = tmp_path / "wb.xlsx"
    write_workbook(path, _result(rows))
    assert openpyxl.load_workbook(path).sheetnames == ["Classified Data", "Summary", "Run Manifest"]


def test_cell_values_are_deterministic(tmp_path):
    rows = [
        _row("JOHN", "MWANGI", "Agikuyu (Kikuyu)", 0.98, "assigned"),
        _row("PETER", "JUMA", "", 0.44, "review"),
    ]
    # xlsxwriter output is not byte-stable (zip mtimes); compare cell values instead.
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    write_workbook(a, _result(rows))
    write_workbook(b, _result(rows))

    def cells(path):
        wb = openpyxl.load_workbook(path)
        return {name: [[c.value for c in row] for row in wb[name].iter_rows()] for name in wb.sheetnames}

    assert cells(a) == cells(b)
