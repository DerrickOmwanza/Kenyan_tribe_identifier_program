"""End-to-end tests for kne.pipeline.classify_file."""

from __future__ import annotations

import csv
import json

import pytest

from kne.pipeline import ClassifyError, classify_file

from conftest import KAMBA, KIKUYU, LUO, make_xlsx


def _rows(path):
    with open(path, encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _classify(inp, out, corpus_db, evidence_dir, test_config, **kw):
    return classify_file(
        inp, out, db_path=corpus_db, reference_dir=evidence_dir,
        config=test_config, **kw,
    )


def test_fresh_mode_appends_columns_and_preserves_input(tmp_path, corpus_db, evidence_dir, test_config):
    src = tmp_path / "in.csv"
    src.write_text(
        "fname,mname,sname,notes\n"
        "JOHN,KAMANDE,MWANGI,keep-me\n"
        "PETER,ODHIAMBO,OCHIENG,two\n"
        "QWER,TYUI,OPAS,nothing\n",
        encoding="utf-8",
    )
    out = tmp_path / "out.csv"
    summary = _classify(src, out, corpus_db, evidence_dir, test_config, mode="fresh")

    rows = _rows(out)
    assert list(rows[0]) == ["fname", "mname", "sname", "notes", "tribe_predicted",
                             "tribe_confidence", "tribe_status", "tribe_basis", "tribe_alt_candidates"]
    assert rows[0]["notes"] == "keep-me"
    assert rows[0]["tribe_predicted"] == KIKUYU and rows[0]["tribe_status"] == "assigned"
    assert rows[2]["tribe_status"] == "review" and rows[2]["tribe_predicted"] == ""
    assert summary["mode"] == "fresh"
    assert (tmp_path / "out.manifest.json").is_file()
    assert (tmp_path / "out.review.csv").is_file()


def test_refine_mode_statuses(tmp_path, corpus_db, evidence_dir, test_config):
    src = tmp_path / "in.csv"
    src.write_text(
        "fname,mname,sname,Tribe\n"
        "JOHN,KAMANDE,MWANGI,Agikuyu (Kikuyu)\n"     # confirmed
        "PETER,ODHIAMBO,OCHIENG,Akamba (Kamba)\n"    # conflict (strongly Luo)
        "QWER,TYUI,OPAS,Luo\n",                       # unverified (no signal)
        encoding="utf-8",
    )
    out = tmp_path / "out.csv"
    _classify(src, out, corpus_db, evidence_dir, test_config, mode="refine")
    rows = _rows(out)
    assert rows[0]["tribe_status"] == "confirmed"
    assert rows[1]["tribe_status"] == "conflict"
    assert rows[1]["tribe_original"] == "Akamba (Kamba)" and rows[1]["tribe_suggested"] == LUO
    assert rows[2]["tribe_status"] == "unverified" and rows[2]["tribe_predicted"] == "Luo"


def test_auto_mode_detects_refine(tmp_path, corpus_db, evidence_dir, test_config):
    src = tmp_path / "in.csv"
    src.write_text("fname,mname,sname,Tribe\nJOHN,KAMANDE,MWANGI,Agikuyu (Kikuyu)\n", encoding="utf-8")
    summary = _classify(src, tmp_path / "o.csv", corpus_db, evidence_dir, test_config, mode="auto")
    assert summary["mode"] == "refine"

    src2 = tmp_path / "in2.csv"
    src2.write_text("fname,mname,sname\nJOHN,KAMANDE,MWANGI\n", encoding="utf-8")
    s2 = _classify(src2, tmp_path / "o2.csv", corpus_db, evidence_dir, test_config, mode="auto")
    assert s2["mode"] == "fresh"


def test_xlsx_input(tmp_path, corpus_db, evidence_dir, test_config):
    src = make_xlsx(tmp_path / "in.xlsx", [("MARY", "MUMBUA", "MUTUA")], header=("fname", "mname", "sname"))
    out = tmp_path / "out.csv"
    _classify(src, out, corpus_db, evidence_dir, test_config, mode="fresh")
    assert _rows(out)[0]["tribe_predicted"] == KAMBA


def test_column_auto_detection(tmp_path, corpus_db, evidence_dir, test_config):
    src = tmp_path / "in.csv"
    src.write_text("First Name,Middle Name,Surname,Ethnicity\nJOHN,KAMANDE,MWANGI,\n", encoding="utf-8")
    out = tmp_path / "out.csv"
    summary = _classify(src, out, corpus_db, evidence_dir, test_config)  # mode=auto
    assert _rows(out)[0]["tribe_predicted"] == KIKUYU
    assert summary["columns_detected"] == {
        "fname": "First Name", "mname": "Middle Name", "sname": "Surname", "tribe": "Ethnicity",
    }


def test_explicit_column_override_missing(tmp_path, corpus_db, evidence_dir, test_config):
    src = tmp_path / "in.csv"
    src.write_text("First,Middle,Surname\nJOHN,KAMANDE,MWANGI\n", encoding="utf-8")
    with pytest.raises(ClassifyError, match="not found"):
        _classify(src, tmp_path / "x.csv", corpus_db, evidence_dir, test_config,
                  mode="fresh", fname_col="given_name")


def test_unresolvable_columns_error(tmp_path, corpus_db, evidence_dir, test_config):
    src = tmp_path / "in.csv"
    src.write_text("a,b,c\n1,2,3\n", encoding="utf-8")
    with pytest.raises(ClassifyError, match="could not find a 'fname' column"):
        _classify(src, tmp_path / "x.csv", corpus_db, evidence_dir, test_config, mode="fresh")


def test_unsupported_output_extension(tmp_path, corpus_db, evidence_dir, test_config):
    src = tmp_path / "in.csv"
    src.write_text("fname,mname,sname\nJOHN,KAMANDE,MWANGI\n", encoding="utf-8")
    with pytest.raises(ClassifyError, match="unsupported output extension"):
        _classify(src, tmp_path / "out.parquet", corpus_db, evidence_dir, test_config, mode="fresh")


def test_xlsx_output(tmp_path, corpus_db, evidence_dir, test_config):
    import openpyxl

    src = tmp_path / "in.csv"
    src.write_text(
        "fname,mname,sname,Tribe\n"
        "JOHN,KAMANDE,MWANGI,Agikuyu (Kikuyu)\n"
        "PETER,ODHIAMBO,OCHIENG,Akamba (Kamba)\n"
        "QWER,TYUI,OPAS,Luo\n",
        encoding="utf-8",
    )
    out = tmp_path / "out.xlsx"
    summary = _classify(src, out, corpus_db, evidence_dir, test_config, mode="refine")
    assert summary["outputs"] == {"workbook": "out.xlsx"}
    assert (tmp_path / "out.manifest.json").is_file()

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Classified Data", "Summary", "Review Queue", "Run Manifest"]
    data = wb["Classified Data"]
    assert data.max_row == 4  # header + 3 rows
    assert data.freeze_panes == "A2"
    header = [c.value for c in data[1]]
    assert header[:4] == ["fname", "mname", "sname", "Tribe"]
    conf_col = header.index("tribe_confidence") + 1
    assert isinstance(data.cell(row=2, column=conf_col).value, float)


def test_xlsx_sheet_selection(tmp_path, corpus_db, evidence_dir, test_config):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    wb.active.append(["ignore"])
    ws = wb.create_sheet("People")
    ws.append(["fname", "mname", "sname"])
    ws.append(["JOHN", "KAMANDE", "MWANGI"])
    src = tmp_path / "multi.xlsx"
    wb.save(src)

    out = tmp_path / "out.csv"
    _classify(src, out, corpus_db, evidence_dir, test_config, mode="fresh", sheet="People")
    assert _rows(out)[0]["tribe_predicted"] == KIKUYU


def test_determinism(tmp_path, corpus_db, evidence_dir, test_config):
    src = tmp_path / "in.csv"
    src.write_text("fname,mname,sname\nJOHN,OCHIENG,OTIENO\nMARY,WANJIRU,WANJIKU\n", encoding="utf-8")
    a, b = tmp_path / "a.csv", tmp_path / "b.csv"
    _classify(src, a, corpus_db, evidence_dir, test_config, mode="fresh")
    _classify(src, b, corpus_db, evidence_dir, test_config, mode="fresh")
    assert a.read_bytes() == b.read_bytes()

    ma = json.loads((tmp_path / "a.manifest.json").read_text())
    mb = json.loads((tmp_path / "b.manifest.json").read_text())
    for m in (ma, mb):
        m.pop("generated_utc")
        m.pop("outputs")
    assert ma == mb


def test_manifest_contents(tmp_path, corpus_db, evidence_dir, test_config):
    src = tmp_path / "in.csv"
    src.write_text("fname,mname,sname\nJOHN,KAMANDE,MWANGI\n", encoding="utf-8")
    _classify(src, tmp_path / "o.csv", corpus_db, evidence_dir, test_config, mode="fresh")
    m = json.loads((tmp_path / "o.manifest.json").read_text())
    assert m["scorer"] == "loglinear:1"
    assert m["target_tribe_count"] == 3  # KIKUYU / LUO / KAMBA in the synthetic corpus
    assert set(m["status_counts"]) <= {"assigned", "review", "generic"}
    assert len(m["corpus_version"]) == 64 and len(m["config_digest"]) == 64
